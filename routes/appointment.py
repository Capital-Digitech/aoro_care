from datetime import datetime

from flask import (
    Blueprint,
    render_template,
    request,
    redirect,
    url_for,
    flash,
    jsonify,
    abort,
    session,
    send_file
)
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

from database import db
from models import Patient, Doctor, Appointment, Hospital


# ==========================================================
# Blueprint
# ==========================================================

appointment_bp = Blueprint(
    "appointment",
    __name__,
    url_prefix="/appointment"
)


# ==========================================================
# Helpers
# ==========================================================

def _parse_date(value):

    if not value:
        return None

    try:
        return datetime.strptime(
            value,
            "%Y-%m-%d"
        ).date()

    except ValueError:
        return None



def _parse_time(value):

    if not value:
        return None

    try:
        return datetime.strptime(
            value,
            "%H:%M"
        ).time()

    except ValueError:
        return None



def _current_doctor():
    # NOTE: this helper was previously defined twice, back to back,
    # with identical bodies (harmless but dead duplicate code). The
    # duplicate has been removed since this function sits directly in
    # the section that powers the My Appointments page.

    if "user" not in session:
        abort(401, description="Login required.")

    if session["user"]["role"] != "doctor":
        abort(403, description="Doctor access required.")

    user_id = session["user"]["id"]

    doctor = Doctor.query.filter_by(
        user_id=user_id
    ).first()

    if doctor is None:
        abort(403, description="Doctor profile not found.")

    return doctor


def _get_doctor_appointment(appointment_id):
    """Fetch an appointment, scoped to the logged-in doctor.

    This is the single authorization choke point for every doctor
    appointment action (view/edit/delete): if the appointment doesn't
    belong to the current doctor, it 404s rather than leaking whether
    the ID exists at all.
    """

    doctor = _current_doctor()

    appointment = Appointment.query.filter_by(
        id=appointment_id,
        doctor_id=doctor.id
    ).first()

    if appointment is None:
        abort(404, description="Appointment not found.")

    return appointment


def _serialize_appointment(appointment):
    patient_name = "-"
    patient_code = "-"

    if appointment.patient and appointment.patient.user:
        patient_name = (
            f"{appointment.patient.user.first_name} "
            f"{appointment.patient.user.last_name}"
        ).strip()
        patient_code = appointment.patient.patient_code or "-"

    doctor_name = "-"

    if appointment.doctor and appointment.doctor.user:
        doctor_name = (
            f"Dr. {appointment.doctor.user.first_name} "
            f"{appointment.doctor.user.last_name}"
        ).strip()

    return dict(
        id=appointment.id,
        patient_name=patient_name,
        patient_code=patient_code,
        doctor_name=doctor_name,
        appointment_date=(
            appointment.appointment_date.strftime("%Y-%m-%d")
            if appointment.appointment_date else ""
        ),
        appointment_time=(
            appointment.appointment_time.strftime("%H:%M")
            if appointment.appointment_time else ""
        ),
        appointment_type=appointment.appointment_type or "",
        status=appointment.status or "scheduled",
        reason=appointment.reason or "",
        meeting_link=appointment.meeting_link or ""
    )


# ==========================================================
# ADMIN APPOINTMENT LIST
# ==========================================================

@appointment_bp.route("/")
def appointment_list():

    appointments = Appointment.query.order_by(
        Appointment.created_at.desc()
    ).all()

    patients = Patient.query.filter_by(
        status="active"
    ).all()

    doctors = Doctor.query.filter_by(
        status="active"
    ).all()

    hospitals = Hospital.query.filter_by(
        status="active"
    ).all()

    return render_template(
        "appointment/appointment_list.html",
        appointments=appointments,
        patients=patients,
        doctors=doctors,
        hospitals=hospitals
    )


# ==========================================================
# ADMIN ADD APPOINTMENT
# ==========================================================

@appointment_bp.route(
    "/add",
    methods=["GET", "POST"]
)
def appointment_add():

    patients = Patient.query.filter_by(
        status="active"
    ).all()

    doctors = Doctor.query.filter_by(
        status="active"
    ).all()

    hospitals = Hospital.query.filter_by(
        status="active"
    ).all()

    if request.method == "POST":

        appointment = Appointment(

            patient_id=request.form.get(
                "patient_id"
            ),

            doctor_id=request.form.get(
                "doctor_id"
            ),

            hospital_id=request.form.get(
                "hospital_id"
            ),

            appointment_date=_parse_date(
                request.form.get(
                    "appointment_date"
                )
            ),

            appointment_time=_parse_time(
                request.form.get(
                    "appointment_time"
                )
            ),

            appointment_type=request.form.get(
                "appointment_type"
            ),

            reason=request.form.get(
                "reason"
            ),

            meeting_link=request.form.get(
                "meeting_link"
            ),

            status=request.form.get(
                "status",
                "scheduled"
            )
        )

        db.session.add(
            appointment
        )

        db.session.commit()

        flash(
            "Appointment added successfully.",
            "success"
        )

        return redirect(
            url_for(
                "appointment.appointment_list"
            )
        )

    return render_template(
        "appointment/add_appointment.html",
        patients=patients,
        doctors=doctors,
        hospitals=hospitals
    )


# ==========================================================
# ADMIN EDIT APPOINTMENT
# ==========================================================

@appointment_bp.route(
    "/edit/<string:id>",
    methods=["GET", "POST"]
)
def appointment_edit(id):

    appointment = Appointment.query.get_or_404(
        id
    )

    patients = Patient.query.filter_by(
        status="active"
    ).all()

    doctors = Doctor.query.filter_by(
        status="active"
    ).all()

    hospitals = Hospital.query.filter_by(
        status="active"
    ).all()

    if request.method == "POST":

        appointment.patient_id = request.form.get(
            "patient_id"
        )

        appointment.doctor_id = request.form.get(
            "doctor_id"
        )

        appointment.hospital_id = request.form.get(
            "hospital_id"
        )

        appointment.appointment_date = _parse_date(
            request.form.get(
                "appointment_date"
            )
        )

        appointment.appointment_time = _parse_time(
            request.form.get(
                "appointment_time"
            )
        )

        appointment.appointment_type = request.form.get(
            "appointment_type"
        )

        appointment.reason = request.form.get(
            "reason"
        )

        appointment.meeting_link = request.form.get(
            "meeting_link"
        )

        appointment.status = request.form.get(
            "status"
        )

        db.session.commit()

        flash(
            "Appointment updated successfully.",
            "success"
        )

        return redirect(
            url_for(
                "appointment.appointment_list"
            )
        )

    return render_template(
        "appointment/edit_appointment.html",
        appointment=appointment,
        patients=patients,
        doctors=doctors,
        hospitals=hospitals
    )
# ==========================================================
# ADMIN EXPORT APPOINTMENTS — EXCEL
# ==========================================================

@appointment_bp.route("/export")
def appointment_export():
    ids = request.args.get("ids", "")

    if not ids:
        return jsonify(
            success=False,
            message="No appointments selected for export."
        ), 400

    appointment_ids = [
        value.strip()
        for value in ids.split(",")
        if value.strip()
    ]

    if not appointment_ids:
        return jsonify(
            success=False,
            message="No appointments selected for export."
        ), 400

    appointments = Appointment.query.filter(
        Appointment.id.in_(appointment_ids)
    ).order_by(
        Appointment.created_at.desc()
    ).all()

    if not appointments:
        return jsonify(
            success=False,
            message="No appointments found."
        ), 404

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Appointments"

    headers = [
        "Patient",
        "Doctor",
        "Appointment Date",
        "Appointment Time",
        "Appointment Type",
        "Status",
        "Reason",
        "Meeting Link"
    ]

    worksheet.append(headers)

    for appointment in appointments:
        patient_name = "-"

        if appointment.patient and appointment.patient.user:
            patient_name = (
                f"{appointment.patient.user.first_name} "
                f"{appointment.patient.user.last_name}"
            ).strip()

        doctor_name = "-"

        if appointment.doctor and appointment.doctor.user:
            doctor_name = (
                f"Dr. {appointment.doctor.user.first_name} "
                f"{appointment.doctor.user.last_name}"
            ).strip()

        appointment_type = (
            appointment.appointment_type or "-"
        ).strip()

        appointment_type_map = {
            "online": "Online",
            "offline": "Offline"
        }

        appointment_type = appointment_type_map.get(
            appointment_type.lower(),
            appointment_type
        )

        status = (appointment.status or "-").strip()

        status_map = {
            "scheduled": "Scheduled",
            "completed": "Completed",
            "cancelled": "Cancelled"
        }

        status = status_map.get(
            status.lower(),
            status
        )

        worksheet.append([
            patient_name,
            doctor_name,
            appointment.appointment_date,
            appointment.appointment_time,
            appointment_type,
            status,
            appointment.reason or "",
            appointment.meeting_link or ""
        ])

    # Header formatting
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # Cell alignment
    for row in worksheet.iter_rows(
        min_row=2,
        max_row=worksheet.max_row
    ):
        for cell in row:
            cell.alignment = Alignment(
                vertical="center"
            )

    # Date/time formatting
    for row in worksheet.iter_rows(
        min_row=2,
        max_row=worksheet.max_row
    ):
        date_cell = row[2]
        time_cell = row[3]

        if date_cell.value:
            date_cell.number_format = "yyyy-mm-dd"

        if time_cell.value:
            time_cell.number_format = "hh:mm AM/PM"

    # Meeting links as clickable hyperlinks
    for row in worksheet.iter_rows(
        min_row=2,
        max_row=worksheet.max_row
    ):
        meeting_cell = row[7]

        if meeting_cell.value:
            meeting_cell.hyperlink = meeting_cell.value
            meeting_cell.style = "Hyperlink"

    # Excel table
    if worksheet.max_row >= 2:
        table_ref = (
            f"A1:H{worksheet.max_row}"
        )

        excel_table = Table(
            displayName="AppointmentsTable",
            ref=table_ref
        )

        table_style = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False
        )

        excel_table.tableStyleInfo = table_style
        worksheet.add_table(excel_table)

    # Column widths
    column_widths = {
        "A": 24,
        "B": 24,
        "C": 18,
        "D": 20,
        "E": 20,
        "F": 16,
        "G": 35,
        "H": 45
    }

    for column, width in column_widths.items():
        worksheet.column_dimensions[column].width = width

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    filename = (
        f"appointments_"
        f"{datetime.now().strftime('%Y-%m-%d')}.xlsx"
    )

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

# ==========================================================
# ADMIN VIEW APPOINTMENT
# ==========================================================

@appointment_bp.route(
    "/view/<string:id>"
)
def appointment_view(id):

    appointment = Appointment.query.get_or_404(
        id
    )

    return render_template(
        "appointment/view_appointment.html",
        appointment=appointment
    )



# ==========================================================
# ADMIN DELETE APPOINTMENT
# ==========================================================

@appointment_bp.route(
    "/delete/<string:id>",
    methods=["POST"]
)
def appointment_delete(id):


    appointment = Appointment.query.get_or_404(
        id
    )


    db.session.delete(
        appointment
    )


    db.session.commit()


    flash(
        "Appointment deleted successfully.",
        "success"
    )


    return redirect(
        url_for(
            "appointment.appointment_list"
        )
    )
# ==========================================================
# DOCTOR APPOINTMENT LIST
# ==========================================================

@appointment_bp.route("/doctor")
def doctor_appointment_list():

    doctor = _current_doctor()

    appointments = Appointment.query.filter_by(
        doctor_id=doctor.id
    ).order_by(
        Appointment.appointment_date.desc()
    ).all()

    return render_template(
        "doctor/doctor_appointment_list.html",
        doctor=doctor,
        appointments=appointments,
        now=datetime.utcnow
    )


# ==========================================================
# DOCTOR VIEW APPOINTMENT (JSON)
#
# Powers the "View" action and also feeds the "Edit" modal's
# pre-fill, so there's one source of truth for what a doctor
# is allowed to see about one of their own appointments.
# ==========================================================

@appointment_bp.route(
    "/doctor/<string:appointment_id>",
    methods=["GET"]
)
def doctor_appointment_detail(appointment_id):

    appointment = _get_doctor_appointment(
        appointment_id
    )

    return jsonify(
        success=True,
        appointment=_serialize_appointment(appointment)
    )


# ==========================================================
# DOCTOR EDIT APPOINTMENT
#
# Replaces the old confirm/complete/cancel/reschedule routes.
# Those set status="confirmed", which is not a valid value in
# the Appointment.status enum (scheduled/completed/cancelled/
# missed) and would fail at commit time — this route validates
# against the real enum instead. A doctor can update date,
# time, type, status and reason for an appointment that
# belongs to them; patient/doctor/hospital assignment stays
# admin-only, unchanged from before.
# ==========================================================

VALID_APPOINTMENT_TYPES = {"online", "offline", "emergency"}
VALID_APPOINTMENT_STATUSES = {"scheduled", "completed", "cancelled", "missed"}


@appointment_bp.route(
    "/doctor/<string:appointment_id>/edit",
    methods=["POST"]
)
def doctor_appointment_edit(appointment_id):

    appointment = _get_doctor_appointment(
        appointment_id
    )

    new_date = _parse_date(
        request.form.get("appointment_date")
    )

    new_time = _parse_time(
        request.form.get("appointment_time")
    )

    appointment_type = request.form.get("appointment_type")
    status = request.form.get("status")
    reason = request.form.get("reason")

    if not new_date or not new_time:
        abort(400, description="Valid date and time required.")

    if appointment_type not in VALID_APPOINTMENT_TYPES:
        abort(400, description="Invalid appointment type.")

    if status not in VALID_APPOINTMENT_STATUSES:
        abort(400, description="Invalid status.")

    appointment.appointment_date = new_date
    appointment.appointment_time = new_time
    appointment.appointment_type = appointment_type
    appointment.status = status
    appointment.reason = reason

    db.session.commit()

    return jsonify(
        success=True,
        message="Appointment updated successfully.",
        appointment=_serialize_appointment(appointment)
    )


# ==========================================================
# DOCTOR DELETE APPOINTMENT
# ==========================================================

@appointment_bp.route(
    "/doctor/<string:appointment_id>/delete",
    methods=["POST"]
)
def doctor_appointment_delete(appointment_id):

    appointment = _get_doctor_appointment(
        appointment_id
    )

    db.session.delete(appointment)
    db.session.commit()

    return jsonify(
        success=True,
        message="Appointment deleted successfully."
    )


# ==========================================================
# ERROR HANDLERS
# ==========================================================

@appointment_bp.errorhandler(400)
def bad_request(error):

    return jsonify(
        success=False,
        message=str(error.description)
    ), 400



@appointment_bp.errorhandler(401)
def unauthorized(error):

    return jsonify(
        success=False,
        message=str(error.description)
    ), 401



@appointment_bp.errorhandler(403)
def forbidden(error):

    return jsonify(
        success=False,
        message=str(error.description)
    ), 403



@appointment_bp.errorhandler(404)
def not_found(error):

    return jsonify(
        success=False,
        message=str(error.description)
    ), 404



@appointment_bp.errorhandler(500)
def internal_error(error):

    db.session.rollback()


    return jsonify(
        success=False,
        message="Internal Server Error."
    ), 500